import contextlib
import io
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from argparse import Namespace
from helpers import d,workspace,deck

class MaterialTests(unittest.TestCase):
    def test_page_ranges(self):
        self.assertEqual(d.parse_pages('1,3-5',5),[1,3,4,5])
        for s in ['0','6','4-2']:
            with self.assertRaises(ValueError):d.parse_pages(s,5)
    def test_path_traversal_rejected(self):
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(ValueError):d.under(Path(t),'../escape')
    def test_pdf_crop_size_and_rotation(self):
        import pypdfium2 as pdfium
        from PIL import Image
        with tempfile.TemporaryDirectory() as t, contextlib.redirect_stdout(io.StringIO()):
            root=Path(t);pdf=root/'source.pdf';doc=pdfium.PdfDocument.new();page=doc.new_page(300,200);page.close();doc.save(str(pdf));doc.close()
            man=root/'figures.json';d.write(man,{'figures':[{'id':'F1','pdf_page':1,'bbox_pt':[20,20,280,100]}]})
            d.crop(Namespace(pdf=pdf,manifest=man,out=root/'crops',dpi=144))
            with Image.open(root/'crops/F1.png') as image:self.assertEqual(image.size,(520,160))
            report=d.load(root/'crops/crop_report.json');self.assertEqual(report[0]['visual_review'],'pending')
            d.write(man,{'figures':[{'id':'F1','pdf_page':1,'bbox_pt':[-1,20,280,100]}]})
            with self.assertRaises(ValueError):d.crop(Namespace(pdf=pdf,manifest=man,out=root/'bad',dpi=144))
    def test_notes_and_hidden_audit(self):
        with tempfile.TemporaryDirectory() as t,contextlib.redirect_stdout(io.StringIO()):
            p=workspace(Path(t));path=deck(p);r=d.audit_pptx(path,p/'design/slide_plan.json')
            self.assertEqual(r['machine_status'],'pass');self.assertEqual(r['hidden_slides'],1);self.assertEqual(r['planned_seconds'],120)
    def test_missing_notes_detected(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'no_notes.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(path)
            self.assertEqual(d.audit_pptx(path)['machine_status'],'fail')
    def test_native_text_inside_nested_groups_is_counted(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'grouped.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
            group=slide.shapes.add_group_shape();nested=group.shapes.add_group_shape()
            nested.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)).text='合成分组文字'
            group.shapes.add_textbox(Inches(1),Inches(2),Inches(4),Inches(1)).text='Synthetic grouped annotation'
            slide.notes_slide.notes_text_frame.text='Synthetic group audit notes.';prs.save(path)
            result=d.audit_pptx(path)
            self.assertEqual(result['machine_status'],'pass')
            self.assertEqual(result['nonempty_native_text_objects'],2)
            self.assertEqual(result['records'][0]['native_text_objects'],2)
            self.assertTrue(any('group-child transformations' in warning for warning in result['warnings']))

@unittest.skipUnless(os.environ.get('RUN_RENDER_TESTS')=='1','Set RUN_RENDER_TESTS=1 for real synthetic rendering tests')
class RenderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not (shutil.which('soffice') or shutil.which('libreoffice')):
            raise RuntimeError('RUN_RENDER_TESTS=1 requires LibreOffice; refusing to skip requested rendering')

    def setUp(self):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup)
        self.root=Path(temp.name)
        # Only synthetic inputs/outputs; keep partial evidence when an assertion fails.
        artifacts=os.environ.get('RENDER_TEST_ARTIFACTS')
        if artifacts:
            self.addCleanup(shutil.copytree,self.root,Path(artifacts)/self._testMethodName,dirs_exist_ok=True)
        sink=contextlib.redirect_stdout(io.StringIO())
        sink.__enter__();self.addCleanup(sink.__exit__,None,None,None)

    def assert_render_output(self,path,out,mapping,includes_hidden,tokens):
        import pypdfium2 as pdfium
        from PIL import Image
        report=d.load(out/'render_report.json')
        self.assertEqual(report['source_pptx_sha256'],d.sha256(path))
        self.assertEqual(report['slide_mapping'],mapping)
        self.assertEqual(report['includes_hidden'],includes_hidden)
        self.assertEqual(report['pdf_pages'],len(mapping))
        self.assertEqual(report['dpi'],96)
        pdf=out/('all_slides.pdf' if includes_hidden else 'main_slides.pdf')
        self.assertEqual(Path(report['pdf']),pdf.resolve())
        self.assertEqual(report['pdf_sha256'],d.sha256(pdf))
        self.assertTrue((out/'renderer.log').is_file())
        manifest=d.load(out/'pages/render_manifest.json')
        self.assertEqual(manifest['renders'],report['renders'])
        self.assertEqual(len(report['renders']),len(mapping))
        expected_files={f'page_{n:03d}.png' for n in range(1,len(mapping)+1)}
        self.assertEqual({p.name for p in (out/'pages').glob('*.png')},expected_files)
        texts=[]
        with pdfium.PdfDocument(str(pdf)) as doc:
            self.assertEqual(len(doc),len(mapping))
            for n,(record,token) in enumerate(zip(report['renders'],tokens),1):
                self.assertEqual(record['pdf_page'],n)
                self.assertEqual(record['file'],f'page_{n:03d}.png')
                image_path=out/'pages'/record['file']
                self.assertEqual(record['sha256'],d.sha256(image_path))
                page=doc[n-1]
                try:
                    width,height=page.get_size()
                    textpage=page.get_textpage()
                    try:text=textpage.get_text_bounded()
                    finally:textpage.close()
                    self.assertIn(token,text)
                    texts.append(text)
                    with Image.open(image_path) as image:
                        self.assertEqual(image.format,'PNG')
                        self.assertAlmostEqual(image.width,width*96/72,delta=1)
                        self.assertAlmostEqual(image.height,height*96/72,delta=1)
                        self.assertTrue(any(low<high for low,high in image.convert('RGB').getextrema()),
                                        'Rendered page must not be a solid/blank image')
                finally:page.close()
        return texts

    def test_actual_renderer_covers_hidden_without_mutating_source(self):
        p=workspace(self.root);path=deck(p);before=d.sha256(path)
        for include_hidden,mapping in [(True,[1,2]),(False,[1])]:
            out=p/'qa'/('all' if include_hidden else 'main')
            d.render_pptx(Namespace(pptx=path,out=out,include_hidden=include_hidden,dpi=96,libreoffice=None))
            self.assert_render_output(path,out,mapping,include_hidden,[f'Synthetic fixture {n}' for n in mapping])
        self.assertEqual(d.sha256(path),before)

    def test_mixed_script_navigation_and_interleaved_backup(self):
        """Native synthetic smoke deck, not ppt-master's application of SVG templates."""
        import copy
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt
        from helpers import w
        p=workspace(self.root,current_design=True)
        style=d.load(p/'design/style_profile.json')
        style['section_navigation']['sections']=[{'id':'example','label':'方法 Method'},
                                                   {'id':'results','label':'结果 Results'}]
        d.write(p/'design/style_profile.json',style)
        plan=d.load(p/'design/slide_plan.json')
        last=copy.deepcopy(plan['slides'][1]);last.update(id='P04',section_id='results',seconds=30)
        plan['slides'][1]['seconds']=75;plan['slides'].append(last)
        d.write(p/'design/slide_plan.json',plan)
        self.assertEqual(w.preflight_report(p)['status'],'pass')
        path=deck(p);prs=Presentation(path)
        labels=['CoverToken 合成封面','MethodToken 中英混排 Agent planning',
                'BackupToken 隐藏备份','ResultToken 结果分析']
        typography=style['typography']

        def text_run(paragraph,text,size):
            run=paragraph.add_run();run.text=text
            run.font.name=typography['latin_font_family'];run.font.size=Pt(size)
            ea=OxmlElement('a:ea');ea.set('typeface',typography['cjk_font_family'])
            run._r.get_or_add_rPr().append(ea)

        for slide,row,label in zip(prs.slides,plan['slides'],labels):
            frame=slide.shapes[0].text_frame;frame.clear()
            text_run(frame.paragraphs[0],label,24)
            if not row['hidden'] and row['id']!='P01':
                for i,section in enumerate(style['section_navigation']['sections']):
                    box=slide.shapes.add_textbox(Inches(1+3*i),Inches(6),Inches(2.5),Inches(.5))
                    text_run(box.text_frame.paragraphs[0],section['label'],14)
                    box.fill.solid()
                    box.fill.fore_color.rgb=RGBColor.from_string('BDD7EE' if section['id']==row['section_id'] else 'EEEEEE')
        prs.save(path);before=d.sha256(path)
        audit=d.audit_pptx(path,p/'design/slide_plan.json');d.write(p/'qa/audit.json',audit)
        self.assertEqual(audit['machine_status'],'pass')
        self.assertEqual(audit['hidden_slides'],1)
        self.assertEqual(audit['visual_status'],'pending')
        for include_hidden,mapping in [(True,[1,2,3,4]),(False,[1,2,4])]:
            out=p/'qa'/('all' if include_hidden else 'main')
            d.render_pptx(Namespace(pptx=path,out=out,include_hidden=include_hidden,dpi=96,libreoffice=None))
            texts=self.assert_render_output(path,out,mapping,include_hidden,[labels[n-1].split()[0] for n in mapping])
            for text,n in zip(texts,mapping):
                # Unicode/text survival, not exact font metrics or glyph appearance.
                self.assertIn(labels[n-1].split()[1],text,
                              'CJK text missing: check font availability (Linux: fonts-noto-cjk) before diagnosing a renderer regression')
                if n in [2,4]:
                    for section in style['section_navigation']['sections']:self.assertIn(section['label'],text)
            if not include_hidden:self.assertNotIn('BackupToken','\n'.join(texts))
        self.assertEqual(d.sha256(path),before)

if __name__=='__main__':unittest.main()
