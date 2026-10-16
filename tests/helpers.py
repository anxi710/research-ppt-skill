"""Synthetic fixtures. No real paper, model result, institution or private slide."""
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/research-ppt/scripts'))
import workflow as w
import decktools as d
from argparse import Namespace


def workspace(root: Path, *, current_design: bool=False) -> Path:
    """Legacy compatibility fixture by default; current_design exercises real defaults."""
    source=root/'synthetic_source.md'
    source.write_text('# Synthetic teaching source\nFor this fixture only: two plus three is five.\n')
    project=root/'talk'
    w.init(Namespace(project=project,source=[source],template=None,minutes=2,language='en',audience='researchers'))
    brief=d.load(project/'brief.json');brief['sources'][0]['version']='synthetic fixture v1'
    if not current_design:brief.pop('design_contract_version')
    brief.update(qa_time_included=False,style_intent='Synthetic neutral style; no institution')
    # Fictional intake record for synthetic tests/examples, never actual user consent.
    brief['intake_confirmation']={'status':'confirmed',
                                  'user_response':'SYNTHETIC FIXTURE ONLY: fictional acceptance of these test requirements.',
                                  'requirements':w.intake_requirements(brief)}
    d.write(project/'brief.json',brief)
    d.write(project/'analysis/evidence.json',{'claims':[{'id':'C1','kind':'illustration','statement':'Synthetic arithmetic demonstration; not an empirical result.','source_ids':['S1'],'locator':'Synthetic source line 2','conditions':[],'limitations':['No scientific interpretation'],'review_status':'verified'}]})
    if current_design:
        style=d.load(project/'design/style_profile.json');style['review_status']='verified'
        style['template_reference']['status']='selected'
        style['section_navigation']['sections']=[{'id':'example','label':'Synthetic example'}]
    else:style={'name':'synthetic','review_status':'verified','basis':'Synthetic legacy fixture; no institution'}
    d.write(project/'design/style_profile.json',style)
    slides=[]
    for i in range(1,3):
        slides.append({'id':f'P{i:02d}','section':'Example','question':'What is the fixture checking?',
                       'topic':'Synthetic review fixture','evidence_ids':['C1'],
                       'primary_visual':{'type':'native_example','asset_ids':[],'description':'Native arithmetic example'},
                       'reading_order':['Read the synthetic label','Read the example'],
                       'on_slide_text':['Synthetic only'],'takeaway':'This is not paper evidence.',
                       'speaker_notes':'Synthetic notes for rendering and gate tests.','seconds':120 if i==1 else 30,'hidden':i==2})
    if current_design:
        import copy
        backup=copy.deepcopy(slides[1]);backup['id']='P03';slides.append(backup)
        slides[0].update(display_title='Synthetic research presentation',on_slide_text=['Synthetic teaching fixture; not empirical research.'],
                         question='What is this synthetic fixture?',topic='Synthetic title information',seconds=15,
                         primary_visual={'type':'information_cover','asset_ids':[],
                                         'description':'Title and supplied synthetic metadata only; no research background.'})
        slides[1].update(hidden=False,seconds=105)
        for slide in slides:
            slide.setdefault('display_title','Synthetic arithmetic teaching example')
            if slide['id']!='P01':slide['section_id']='example'
    d.write(project/'design/slide_plan.json',{'slides':slides})
    return project


def deck(project: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation()
    plan=d.load(project/'design/slide_plan.json')['slides']
    for i,slide in enumerate(plan):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text=f'Synthetic fixture {i+1}: 2 + 3 = 5'
        sl.notes_slide.notes_text_frame.text='Synthetic only; notes used for structural tests.'
        if slide['hidden']:sl._element.set('show','0')
    path=project/'exports/synthetic.pptx';prs.save(path);return path


def mock_review_fixture(project: Path) -> Namespace:
    """Deliberately fabricated UNIT-TEST records; never used as an actual deck approval."""
    from PIL import Image
    pptx=deck(project);w.handoff(Namespace(project=project,run_id='unit'))
    handoff=project/'handoff/unit';out=project/'qa/unit';(out/'pages').mkdir(parents=True)
    renders=[]
    count=len(d.load(project/'design/slide_plan.json')['slides'])
    for i in range(1,count+1):
        path=out/'pages'/f'page_{i:03d}.png';Image.new('RGB',(100,50),(255,255,255)).save(path)
        renders.append({'pdf_page':i,'file':path.name,'sha256':d.sha256(path)})
    render_report=out/'render_report.json'
    d.write(render_report,{'source_pptx_sha256':d.sha256(pptx),'includes_hidden':True,'slide_mapping':list(range(1,count+1)),'renders':renders})
    review=out/'review.json';w.review_init(Namespace(project=project,pptx=pptx,handoff=handoff,render_report=render_report,out=review))
    r=d.load(review);r['reviewer']='UNIT TEST - not a real reviewer';r['review_mode']='human'
    for row in r['slides']:
        row.update({c:'pass' for c in w.CHECKS});row['comments']='Synthetic UNIT-TEST approval only.'
    d.write(review,r)
    audit=out/'audit.json';d.write(audit,d.audit_pptx(pptx,project/'design/slide_plan.json'))
    receipt=out/'receipt.json'
    d.write(receipt,{'status':'completed','backend_name':'synthetic_test_backend','backend_path':'test fixture','backend_revision':'fixture-1','commands_executed':['Synthetic test construction only'],'documents_read':['tests/helpers.py'],'handoff_manifest_sha256':d.sha256(handoff/'manifest.json'),'pptx_sha256':d.sha256(pptx)})
    return Namespace(project=project,pptx=pptx,handoff=handoff,render_report=render_report,review=review,audit=audit,receipt=receipt,out=out/'gate.json')
