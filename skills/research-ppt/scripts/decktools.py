#!/usr/bin/env python3
"""Research PPT material and artifact helpers. No LLM/API calls; semantic and visual review remain agent tasks.

Python 3.11. All output paths are explicit. No source is modified in place. Production is delegated to ppt-master.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import math
import posixpath
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from typing import Any


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda: f.read(1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding='utf-8'))


def write(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def under(root: Path, relative: str) -> Path:
    result = (root / relative).resolve()
    if not result.is_relative_to(root.resolve()):
        raise ValueError(f'Path leaves project: {relative}')
    return result


def parse_pages(expr: str, count: int) -> list[int]:
    if expr == 'all':
        return list(range(1, count + 1))
    values: set[int] = set()
    for block in expr.split(','):
        bounds = [int(v) for v in block.split('-')]
        if len(bounds) == 1:
            values.add(bounds[0])
        elif len(bounds) == 2 and bounds[0] <= bounds[1]:
            values.update(range(bounds[0], bounds[1] + 1))
        else:
            raise ValueError(f'Invalid page range: {block}')
    if not values or min(values) < 1 or max(values) > count:
        raise ValueError(f'Pages must be in 1..{count}')
    return sorted(values)


def inspect_template(args: argparse.Namespace) -> int:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.dml import MSO_COLOR_TYPE
    prs = Presentation(args.pptx)
    out = args.out.resolve(); out.mkdir(parents=True, exist_ok=True)
    fonts: Counter = Counter(); sizes: Counter = Counter(); colors: Counter = Counter()
    rows = []; assets = {}; asset_dir = out / 'images'; asset_dir.mkdir(exist_ok=True)

    def visit(shapes, slide: int) -> list[dict]:
        info = []
        for sh in shapes:
            entry: dict = {'name': sh.name, 'shape_id': sh.shape_id,
                           'bbox_in': [round(float(v) / 914400, 4) for v in
                                       [sh.left, sh.top, sh.width, sh.height]]}
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                entry['children'] = visit(sh.shapes, slide)
            if sh.has_text_frame:
                entry['text'] = sh.text
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.name: fonts[r.font.name] += 1
                        if r.font.size: sizes[str(round(r.font.size.pt, 2))] += 1
                        if r.font.color.type == MSO_COLOR_TYPE.RGB:
                            colors[str(r.font.color.rgb)] += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = sh.image.blob; digest = hashlib.sha256(blob).hexdigest()
                if digest not in assets:
                    filename = f'image_{len(assets)+1:03d}.{sh.image.ext}'
                    (asset_dir / filename).write_bytes(blob)
                    assets[digest] = {'path': f'images/{filename}', 'sha256': digest,
                                      'size': list(sh.image.size)}
                entry['image'] = assets[digest]['path']
            info.append(entry)
        return info

    for n, slide in enumerate(prs.slides, 1):
        rows.append({'slide': n, 'hidden': slide._element.get('show') == '0',
                     'shapes': visit(slide.shapes, n)})
    report = {'source': str(args.pptx.resolve()), 'sha256': sha256(args.pptx),
              'width_in': prs.slide_width / 914400, 'height_in': prs.slide_height / 914400,
              'slides': rows, 'explicit_fonts': dict(fonts), 'explicit_font_sizes_pt': dict(sizes),
              'explicit_text_colors': dict(colors), 'images': list(assets.values()),
              'limits': ['Counts cover explicit native runs, not all inherited master/theme formatting.',
                         'Image identity, official brand colors, and intended layouts require visual review.',
                         'Group child boxes may be in group-local coordinates.']}
    write(out / 'template_inventory.json', report)
    print(out / 'template_inventory.json')
    return 0



def render_pdf_file(pdf: Path, out: Path, pages: str = 'all', dpi: int = 160) -> dict:
    import pypdfium2 as pdfium
    if not 72 <= dpi <= 600:
        raise ValueError('DPI must be in 72..600')
    out.mkdir(parents=True, exist_ok=True)
    records = []
    with pdfium.PdfDocument(str(pdf)) as doc:
        count = len(doc)
        for n in parse_pages(pages, count):
            page = doc[n - 1]
            bitmap = page.render(scale=dpi / 72)
            image = bitmap.to_pil()
            filename = f'page_{n:03d}.png'
            image.save(out / filename)
            image.close(); bitmap.close(); page.close()
            records.append({'pdf_page': n, 'file': filename, 'sha256': sha256(out / filename)})
    report = {'pdf': str(pdf.resolve()), 'pdf_sha256': sha256(pdf), 'pdf_pages': count,
              'renderer': 'pypdfium2/PDFium', 'dpi': dpi, 'renders': records}
    write(out / 'render_manifest.json', report)
    return report


def render_pdf(args: argparse.Namespace) -> int:
    report = render_pdf_file(args.pdf, args.out, args.pages, args.dpi)
    print(f'Rendered {len(report["renders"])} pages into {args.out}')
    return 0


def extract_pdf(args: argparse.Namespace) -> int:
    import pypdfium2 as pdfium
    result = []
    with pdfium.PdfDocument(str(args.pdf)) as doc:
        for n in parse_pages(args.pages, len(doc)):
            page = doc[n-1]; textpage = page.get_textpage()
            result.append(f'## PDF page {n}\n\n' + textpage.get_text_bounded())
            textpage.close(); page.close()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text('\n\n'.join(result)+'\n', encoding='utf-8')
    print(f'{args.out}: reading aids only; inspect figure/table page images separately.')
    return 0


def crop(args: argparse.Namespace) -> int:
    import pypdfium2 as pdfium
    items = load(args.manifest)
    if isinstance(items, dict): items = items['figures']
    if not isinstance(items, list): raise ValueError('Manifest must contain a figure list')
    args.out.mkdir(parents=True, exist_ok=True)
    report = []; names = set(); digest = sha256(args.pdf)
    with pdfium.PdfDocument(str(args.pdf)) as doc:
        for item in items:
            name = item.get('id')
            if not name or any(c not in 'abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-' for c in name):
                raise ValueError(f'Invalid asset id: {name}')
            if name in names: raise ValueError(f'Duplicate asset id: {name}')
            names.add(name)
            page_no = item.get('pdf_page')
            if type(page_no) is not int or not 1 <= page_no <= len(doc):
                raise ValueError(f'Invalid PDF page for {name}')
            page = doc[page_no - 1]
            try:
                if page.get_rotation():
                    raise ValueError('Rotated page: normalize a separate copy and re-check coordinates first.')
                width, height = page.get_size()
                box = item.get('bbox_pt')
                if not isinstance(box, list) or len(box) != 4 or not all(type(v) in (int,float) and math.isfinite(v) for v in box):
                    raise ValueError(f'Invalid bbox_pt for {name}')
                x0,y0,x1,y1 = box
                if not (0 <= x0 < x1 <= width and 0 <= y0 < y1 <= height):
                    raise ValueError(f'Crop outside page: {name} {box}; page {width}x{height}')
                dpi = item.get('dpi', args.dpi)
                if type(dpi) not in (int,float) or not 72 <= dpi <= 600:
                    raise ValueError('Crop DPI must be in 72..600')
                path = args.out / f'{name}.png'
                # Manifest: top-left box. PDFium: amounts removed at left,bottom,right,top.
                bitmap = page.render(scale=dpi/72, crop=(x0,height-y1,width-x1,y0))
                image = bitmap.to_pil(); pixels = list(image.size); image.save(path)
                image.close(); bitmap.close()
                report.append({**item, 'file': path.name, 'pdf_sha256': digest,
                               'dpi': dpi, 'pixels': pixels, 'sha256': sha256(path),
                               'operation': 'original_pdf_crop', 'visual_review': 'pending'})
            finally:
                page.close()
    write(args.out / 'crop_report.json', report)
    print(f'Created {len(report)} crops. Axis/legend completeness still requires review.')
    return 0
def montage(args: argparse.Namespace) -> int:
    from PIL import Image, ImageDraw
    import re
    def key(p): return [int(s) if s.isdigit() else s for s in re.split(r'(\d+)', p.name)]
    images = sorted([p for p in args.images.iterdir() if p.suffix.lower() in ['.png','.jpg','.jpeg']
                     and p.resolve() != args.out.resolve()], key=key)
    if not images: raise ValueError('No images found')
    cols = args.columns
    if cols < 1 or cols > 10: raise ValueError('columns must be in 1..10')
    tw, th = 336, 220
    canvas = Image.new('RGB', (cols * tw, math.ceil(len(images)/cols) * th), '#E8ECF1')
    for n, path in enumerate(images):
        with Image.open(path) as raw:
            im = raw.convert('RGB'); im.thumbnail((320,190))
        x, y = (n % cols) * tw, (n // cols) * th
        canvas.paste(im, (x + (tw-im.width)//2, y+6+(190-im.height)//2))
        ImageDraw.Draw(canvas).text((x+9, y+201), f'{n+1:02d}  {path.name.encode("ascii","replace").decode()}', fill='#202F42')
    args.out.parent.mkdir(parents=True, exist_ok=True); canvas.save(args.out)
    print(args.out)
    return 0



def render_pptx(args: argparse.Namespace) -> int:
    from pptx import Presentation
    from lxml import etree
    lo = args.libreoffice or shutil.which('libreoffice') or shutil.which('soffice')
    if not lo:
        raise RuntimeError('Real PPTX renderer unavailable. Do not substitute SVG previews for PPTX QA.')
    out = args.out.resolve()
    if out.exists() and any(out.iterdir()):
        raise ValueError('Use a new render directory for each artifact revision.')
    out.mkdir(parents=True, exist_ok=True)
    source_digest = sha256(args.pptx)
    prs = Presentation(args.pptx)
    mapping = [i for i,s in enumerate(prs.slides,1) if args.include_hidden or s._element.get('show') != '0']
    slide_parts = {str(s.part.partname).lstrip('/') for s in prs.slides}
    if not mapping: raise ValueError('No slides selected for rendering')
    with tempfile.TemporaryDirectory(prefix='research_ppt_render_') as temp:
        root = Path(temp); local = root / 'deck.pptx'
        # Preserve all package parts. Only toggle visibility in a temporary copy for QA.
        # Never save the production deck via python-pptx to change this flag.
        if args.include_hidden:
            parser = etree.XMLParser(resolve_entities=False,no_network=True)
            with ZipFile(args.pptx) as src, ZipFile(local,'w',ZIP_DEFLATED) as dest:
                for info in src.infolist():
                    data = src.read(info.filename)
                    if info.filename in slide_parts:
                        tree = etree.fromstring(data,parser); tree.set('show','1')
                        data = etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
                    dest.writestr(info,data)
        else:
            shutil.copy2(args.pptx,local)
        result = subprocess.run([str(lo), '-env:UserInstallation='+(root/'lo_profile').as_uri(),
                                 '--headless','--convert-to','pdf','--outdir',str(root),str(local)],
                                capture_output=True,text=True,timeout=180)
        (out/'renderer.log').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
        pdf = root/'deck.pdf'
        if result.returncode or not pdf.is_file(): raise RuntimeError('PPTX render failed; see renderer.log')
        dest = out/('all_slides.pdf' if args.include_hidden else 'main_slides.pdf')
        shutil.copy2(pdf,dest)
        report = render_pdf_file(dest,out/'pages','all',args.dpi)
        if report['pdf_pages'] != len(mapping):
            raise RuntimeError('Rendered page count differs from selected PPTX slide count')
        report.update({'source_pptx':str(args.pptx.resolve()),'source_pptx_sha256':source_digest,
                       'pptx_renderer':str(lo),'includes_hidden':args.include_hidden,
                       'slide_mapping':mapping})
        if sha256(args.pptx) != source_digest: raise RuntimeError('Source changed during rendering')
        write(out/'render_report.json',report)
    print(out/'render_report.json')
    return 0
def audit_pptx(pptx: Path, plan: Path | None = None) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from lxml import etree
    errors: list[str] = []; warnings: list[str] = []; external = []
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(pptx) as z:
        if z.testzip(): errors.append('ZIP CRC failure')
        files = set(z.namelist())
        for name in files:
            if name.endswith(('.xml','.rels')):
                try: tree = etree.fromstring(z.read(name), parser)
                except etree.XMLSyntaxError as e:
                    errors.append(f'Invalid XML {name}: {e}'); continue
                if name.endswith('.rels'):
                    base = posixpath.dirname(posixpath.dirname(name))
                    for rel in tree:
                        if rel.get('TargetMode') == 'External':
                            external.append(rel.get('Target')); continue
                        target = rel.get('Target','').split('#',1)[0]
                        if not target: continue
                        part = posixpath.normpath(posixpath.join(base,target)).lstrip('/')
                        if part not in files: errors.append(f'Unresolved relationship {name} -> {target}')
    def shape_tree(shapes, nested=False):
        for shape in shapes:
            yield shape,nested
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shape_tree(shape.shapes,True)

    prs = Presentation(pptx); records = []; text_count = 0
    for i, sl in enumerate(prs.slides, 1):
        if not sl.has_notes_slide or not sl.notes_slide.notes_text_frame.text.strip():
            errors.append(f'Slide {i}: missing speaker notes')
        texts = []; images = 0
        for sh,nested in shape_tree(sl.shapes):
            if sh.width < 0 or sh.height < 0: errors.append(f'Slide {i}: negative shape extent')
            # Children use group-local coordinates. Count them, but do not mistake
            # those coordinates for slide-space bounds without applying transforms.
            if not nested and (min(sh.left,sh.top) < -19050 or sh.left+sh.width > prs.slide_width+19050 or sh.top+sh.height > prs.slide_height+19050):
                warnings.append(f'Slide {i}: out-of-canvas box {sh.name}; inspect content/intentional bleed')
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                warnings.append(f'Slide {i}: group-child transformations require visual check')
            if sh.rotation: warnings.append(f'Slide {i}: rotated object requires visual check')
            if sh.has_text_frame and sh.text.strip(): texts.append(sh.text); text_count += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: images += 1
        records.append({'slide': i, 'hidden': sl._element.get('show') == '0',
                        'native_text_objects': len(texts), 'picture_objects': images})
    planned_seconds = None
    if plan:
        data = load(plan); rows = data['slides'] if isinstance(data,dict) else data
        if len(rows) != len(records): errors.append('Plan/PPTX slide count mismatch')
        if len(rows) == len(records):
            for actual, planned in zip(records, rows):
                if actual['hidden'] != bool(planned.get('hidden',False)):
                    errors.append(f'Slide {actual["slide"]}: plan/PPTX hidden status mismatch')
        durations = [r.get('seconds') for r in rows if not r.get('hidden',False)]
        if any(not isinstance(v,(int,float)) or isinstance(v,bool) or not math.isfinite(v) or v <= 0 for v in durations):
            errors.append('Visible slide timing must be a positive finite number')
        else: planned_seconds = sum(durations)
    return {'pptx': str(pptx.resolve()), 'pptx_sha256': sha256(pptx), 'slides': len(records),
            'plan_sha256': sha256(plan) if plan else None,
            'visible_slides': sum(not r['hidden'] for r in records), 'hidden_slides': sum(r['hidden'] for r in records),
            'nonempty_native_text_objects': text_count, 'planned_seconds': planned_seconds,
            'records': records, 'errors': errors, 'warnings': warnings, 'external_links': external,
            'machine_status': 'fail' if errors else 'pass', 'visual_status': 'pending',
            'limits': ['Box checks do not establish absence of text overflow, overlaps, or clipped legends.',
                       'Does not judge claims, chart values, scientific meaning, or font substitution.',
                       'Rendered-page review is mandatory and must match the final PPTX hash.']}


def audit(args: argparse.Namespace) -> int:
    report = audit_pptx(args.pptx, args.plan); write(args.out, report)
    print(f'{report["machine_status"]}: {report["slides"]} slides; visual review pending -> {args.out}')
    return 1 if report['errors'] else 0



def parser() -> argparse.ArgumentParser:
    p=argparse.ArgumentParser(description=__doc__); sub=p.add_subparsers(dest='cmd',required=True)
    s=sub.add_parser('inspect-template');s.add_argument('--pptx',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.set_defaults(func=inspect_template)
    s=sub.add_parser('extract-pdf');s.add_argument('--pdf',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.add_argument('--pages',default='all');s.set_defaults(func=extract_pdf)
    s=sub.add_parser('render-pdf');s.add_argument('--pdf',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.add_argument('--pages',default='all');s.add_argument('--dpi',type=int,default=160);s.set_defaults(func=render_pdf)
    s=sub.add_parser('crop');s.add_argument('--pdf',type=Path,required=True);s.add_argument('--manifest',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.add_argument('--dpi',type=int,default=360);s.set_defaults(func=crop)
    s=sub.add_parser('montage');s.add_argument('--images',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.add_argument('--columns',type=int,default=4);s.set_defaults(func=montage)
    s=sub.add_parser('render-pptx');s.add_argument('--pptx',type=Path,required=True);s.add_argument('--out',type=Path,required=True);s.add_argument('--include-hidden',action='store_true');s.add_argument('--dpi',type=int,default=150);s.add_argument('--libreoffice');s.set_defaults(func=render_pptx)
    s=sub.add_parser('audit');s.add_argument('--pptx',type=Path,required=True);s.add_argument('--plan',type=Path);s.add_argument('--out',type=Path,required=True);s.set_defaults(func=audit)
    return p

if __name__=='__main__':
    try:
        args=parser().parse_args();raise SystemExit(args.func(args))
    except (ValueError,OSError,KeyError,RuntimeError,ImportError,subprocess.SubprocessError) as exc:
        print(f'ERROR: {exc}',file=sys.stderr);raise SystemExit(2)
